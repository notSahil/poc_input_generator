"""Unit tests for Sitetracker Data Hub User Guide generator script."""

from pathlib import Path
import pytest
from docx import Document
from scripts.generate_user_guide_docx import UserGuideDocxBuilder, OUTPUT_DOCX_PATH


def test_user_guide_docx_generation(tmp_path: Path):
    """Verify that UserGuideDocxBuilder generates a valid .docx with all chapters, tables, and placeholders."""
    # Arrange
    target_file = tmp_path / "test_user_guide.docx"
    builder = UserGuideDocxBuilder(target_file)

    # Act
    result_path = builder.build()

    # Assert
    assert result_path.exists()
    assert result_path.stat().st_size > 10000

    doc = Document(str(result_path))

    # Collect text from both root paragraphs and all table cells
    para_text = "\n".join([p.text for p in doc.paragraphs])
    table_text = "\n".join([
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    ])
    full_text = para_text + "\n" + table_text

    # Check key chapters and invariants
    assert "Sitetracker Data Hub" in full_text
    assert "Platform Overview & Business Impact" in full_text
    assert "The 5 Standard Output Files" in full_text
    assert "final_input_file.csv" in full_text
    assert "field_level_changes.csv" in full_text
    assert "invalid_primary_key.csv" in full_text
    assert "duplicate_primary_keys.csv" in full_text
    assert "run_summary.txt" in full_text
    assert "Access & Authentication Gateway" in full_text
    assert "Guided Report Pipeline" in full_text
    assert "Ad-Hoc Object Ingestion & Mapping" in full_text
    assert "Run History & Audit Replay" in full_text
    assert "Interactive Schema & Mapping Editor" in full_text
    assert "Task Scheduler & Persistent Job Queue" in full_text
    assert "Real-Time Telemetry & Live Monitor" in full_text
    assert "Troubleshooting & Quick FAQ" in full_text
    assert "Production Server & DevOps Reference" in full_text

    # Check that screenshot placeholder instructions and figures are present
    assert "📷 [PASTE SCREENSHOT HERE]" in full_text
    assert "Figure 3.1" in full_text
    assert "Figure 5.1" in full_text
    assert "Figure 6.1" in full_text


def test_production_output_doc_exists():
    """Verify the production user guide document exists in docs/ directory."""
    assert OUTPUT_DOCX_PATH.exists()
    assert OUTPUT_DOCX_PATH.stat().st_size > 20000
